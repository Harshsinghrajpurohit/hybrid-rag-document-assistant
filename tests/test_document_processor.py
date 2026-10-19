"""
Tests for the multi-format document processor (Phase 6-C).

All files are built in pytest's tmp_path; no database or Ollama server is
required, and the real backend/chroma_db is never touched.
"""

import pytest

from backend.app.exceptions import ExtractionFailedError
from backend.app.services import document_processor
from backend.app.services.document_processor import process_document


def make_pdf_bytes(text: str = "Hybrid RAG sample page about retrieval.") -> bytes:
    """Build a one-page PDF with extractable text (same trick as test_api)."""
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx_bytes(paragraphs) -> bytes:
    """Build a DOCX in memory. paragraphs: list of (text, style) tuples."""
    import docx

    document = docx.Document()
    for text, style in paragraphs:
        if style:
            document.add_paragraph(text, style=style)
        else:
            document.add_paragraph(text)
    import io

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx_bytes(sheets) -> bytes:
    """Build an XLSX in memory. sheets: {name: (header, rows)}."""
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, (header, rows) in sheets.items():
        sheet = workbook.create_sheet(title=name)
        sheet.append(header)
        for row in rows:
            sheet.append(row)
    import io

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ============================================================================
# PDF regression (existing behavior must be preserved)
# ============================================================================


def test_pdf_regression(tmp_path):
    path = tmp_path / "sample.pdf"
    path.write_bytes(make_pdf_bytes())
    chunks = process_document(str(path), "sample.pdf")

    assert chunks, "PDF should produce chunks"
    assert all(c["file_type"] == "pdf" for c in chunks)
    assert all(isinstance(c["page"], int) and c["page"] >= 1 for c in chunks)
    assert [c["chunk_index"] for c in chunks] == list(range(len(chunks)))
    assert "retrieval" in chunks[0]["text"]


# ============================================================================
# TXT
# ============================================================================


def test_txt_plain(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Plain text notes about hybrid search.", encoding="utf-8")
    chunks = process_document(str(path), "notes.txt")

    assert len(chunks) == 1
    assert chunks[0]["file_type"] == "txt"
    assert "hybrid search" in chunks[0]["text"]
    # No fabricated page or location metadata for TXT
    assert "page" not in chunks[0]
    assert "section" not in chunks[0]


def test_txt_latin1_fallback(tmp_path):
    path = tmp_path / "latin.txt"
    path.write_bytes(b"caf\xe9 notes with latin-1 bytes")
    chunks = process_document(str(path), "latin.txt")

    assert chunks[0]["file_type"] == "txt"
    assert "caf" in chunks[0]["text"]


def test_txt_empty_raises(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_bytes(b"   \n\n  ")
    with pytest.raises(ExtractionFailedError):
        process_document(str(path), "empty.txt")


# ============================================================================
# Markdown
# ============================================================================


def test_md_heading_sections(tmp_path):
    path = tmp_path / "guide.md"
    path.write_text(
        "# RAG Guide\nIntro paragraph about retrieval.\n\n"
        "## Setup\nInstall the backend and start Ollama.\n",
        encoding="utf-8",
    )
    chunks = process_document(str(path), "guide.md")

    assert all(c["file_type"] == "md" for c in chunks)
    intro = next(c for c in chunks if "Intro paragraph" in c["text"])
    setup = next(c for c in chunks if "Install the backend" in c["text"])
    assert intro["section"] == "RAG Guide"
    assert setup["section"] == "RAG Guide > Setup"


def test_md_code_fence_headings_ignored(tmp_path):
    path = tmp_path / "code.md"
    path.write_text(
        "# Top\n```\n# not a real heading\n```\nsome text\n", encoding="utf-8"
    )
    chunks = process_document(str(path), "code.md")

    assert all(c["section"] == "Top" for c in chunks)
    assert any("# not a real heading" in c["text"] for c in chunks)


# ============================================================================
# DOCX
# ============================================================================


def test_docx_sections(tmp_path):
    path = tmp_path / "report.docx"
    path.write_bytes(
        make_docx_bytes(
            [
                ("Intro text before any heading.", None),
                ("Setup", "Heading 1"),
                ("Install steps live here.", "Normal"),
            ]
        )
    )
    chunks = process_document(str(path), "report.docx")

    assert all(c["file_type"] == "docx" for c in chunks)
    intro = next(c for c in chunks if "Intro text" in c["text"])
    setup = next(c for c in chunks if "Install steps" in c["text"])
    assert "section" not in intro  # text before first heading has no section
    assert setup["section"] == "Setup"


def test_docx_malformed_raises(tmp_path):
    path = tmp_path / "broken.docx"
    path.write_bytes(b"this is not a zip archive")
    with pytest.raises(ExtractionFailedError):
        process_document(str(path), "broken.docx")


# ============================================================================
# CSV
# ============================================================================


def test_csv_header_aware_rows(tmp_path):
    path = tmp_path / "data.csv"
    path.write_text("Region,Revenue\nEMEA,4200\nAPAC,5100\n", encoding="utf-8")
    chunks = process_document(str(path), "data.csv")

    assert all(c["file_type"] == "csv" for c in chunks)
    text = "\n".join(c["text"] for c in chunks)
    assert "Region: EMEA" in text
    assert "Revenue: 4200" in text
    assert chunks[0]["rows"] == "2-3"
    assert "sheet" not in chunks[0]  # CSVs have no sheets
    assert "page" not in chunks[0]


def test_csv_header_only_raises(tmp_path):
    path = tmp_path / "header_only.csv"
    path.write_text("A,B\n", encoding="utf-8")
    with pytest.raises(ExtractionFailedError):
        process_document(str(path), "header_only.csv")


def test_csv_row_limit_enforced(tmp_path, monkeypatch):
    monkeypatch.setattr(document_processor, "MAX_SPREADSHEET_ROWS_PER_SHEET", 5)
    path = tmp_path / "big.csv"
    rows = "\n".join(f"id{i},value{i}" for i in range(6))
    path.write_text(f"id,value\n{rows}\n", encoding="utf-8")

    with pytest.raises(ExtractionFailedError) as exc_info:
        process_document(str(path), "big.csv")
    # Explicit error, never silent truncation
    assert "exceeding the limit" in exc_info.value.detail


# ============================================================================
# XLSX
# ============================================================================


def test_xlsx_sheets_and_rows(tmp_path):
    path = tmp_path / "book.xlsx"
    path.write_bytes(
        make_xlsx_bytes(
            {
                "Sales": (["Item", "Revenue"], [["Widget", 120], ["Gadget", 340]]),
                "Targets": (["Item", "Goal"], [["Widget", 200]]),
            }
        )
    )
    chunks = process_document(str(path), "book.xlsx")

    assert all(c["file_type"] == "xlsx" for c in chunks)
    sales = [c for c in chunks if c["sheet"] == "Sales"]
    targets = [c for c in chunks if c["sheet"] == "Targets"]
    assert sales and targets
    assert "Item: Widget | Revenue: 120" in sales[0]["text"]
    assert sales[0]["rows"] == "2-3"
    assert targets[0]["rows"] == "2-2"
    assert "page" not in chunks[0]


def test_xlsx_row_limit_enforced(tmp_path, monkeypatch):
    monkeypatch.setattr(document_processor, "MAX_SPREADSHEET_ROWS_PER_SHEET", 3)
    path = tmp_path / "big.xlsx"
    path.write_bytes(
        make_xlsx_bytes({"Big": (["n"], [[i] for i in range(4)])})
    )

    with pytest.raises(ExtractionFailedError) as exc_info:
        process_document(str(path), "big.xlsx")
    assert "exceeding the limit" in exc_info.value.detail
    assert "Big" in exc_info.value.detail


# ============================================================================
# Dispatcher
# ============================================================================


def test_unsupported_extension_raises(tmp_path):
    path = tmp_path / "slides.pptx"
    path.write_bytes(b"irrelevant")
    with pytest.raises(ExtractionFailedError):
        process_document(str(path), "slides.pptx")