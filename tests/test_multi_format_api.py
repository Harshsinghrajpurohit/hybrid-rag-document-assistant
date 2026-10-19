"""
API-level tests for multi-format document upload (Phase 6-C).

Reuses the isolated temp-ChromaDB pattern from test_api.py and fully mocks
the Ollama HTTP layer, so no Ollama server is required. The real
backend/chroma_db is never touched.
"""

import io
import shutil
import tempfile
from unittest.mock import Mock, patch

import pytest
from fastapi.testclient import TestClient

from backend.app.config import Config
from backend.app.main import app


def make_pdf_bytes(text: str = "Hybrid RAG test document for API upload.") -> bytes:
    import fitz

    doc = fitz.open()
    page = doc.new_page()
    page.insert_text((72, 72), text)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx_bytes(paragraphs) -> bytes:
    import docx

    document = docx.Document()
    for text, style in paragraphs:
        if style:
            document.add_paragraph(text, style=style)
        else:
            document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_xlsx_bytes(sheets) -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for name, (header, rows) in sheets.items():
        sheet = workbook.create_sheet(title=name)
        sheet.append(header)
        for row in rows:
            sheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


SAMPLES = {
    "doc.pdf": lambda: make_pdf_bytes(),
    "report.docx": lambda: make_docx_bytes(
        [("Installation", "Heading 1"), ("Install the backend first.", "Normal")]
    ),
    "notes.txt": lambda: b"Plain text notes about plainwords in documents.",
    "guide.md": lambda: b"# Markdown Guide\nAll about markdown structure.\n",
    "data.csv": lambda: b"Region,Revenue\nEMEA,4200\nAPAC,5100\n",
    "book.xlsx": lambda: make_xlsx_bytes(
        {"Sales": (["Item", "Revenue"], [["Widget", 120], ["Gadget", 340]])}
    ),
}


@pytest.fixture
def mock_embed():
    """Deterministic embeddings for any text - no Ollama required."""
    mock_resp = Mock()
    mock_resp.json.return_value = {"embeddings": [[0.01] * 768]}
    mock_resp.raise_for_status = Mock()
    with patch(
        "backend.app.services.embeddings.requests.post", return_value=mock_resp
    ) as m:
        yield m


@pytest.fixture
def client():
    """TestClient with an isolated temporary ChromaDB directory."""
    original = Config.CHROMA_DB_PATH
    temp_db = tempfile.mkdtemp(prefix="mf_api_")
    Config.CHROMA_DB_PATH = temp_db
    yield TestClient(app)
    Config.CHROMA_DB_PATH = original
    shutil.rmtree(temp_db, ignore_errors=True)


def upload(client, filename: str, content: bytes):
    return client.post(
        "/documents/upload",
        files={"file": (filename, io.BytesIO(content), "application/octet-stream")},
    )


@pytest.mark.parametrize("filename", list(SAMPLES.keys()))
def test_upload_each_format(client, mock_embed, filename):
    response = upload(client, filename, SAMPLES[filename]())

    assert response.status_code == 200, response.text
    data = response.json()
    assert data["file_name"] == filename
    assert data["file_type"] == filename.rsplit(".", 1)[-1]
    assert data["num_chunks"] >= 1
    assert data["doc_id"]


def test_upload_unsupported_extension(client, mock_embed):
    response = upload(client, "slides.pptx", b"pretend presentation")

    assert response.status_code == 400
    assert response.json()["error"] == "UNSUPPORTED_FORMAT"


def test_upload_empty_file(client, mock_embed):
    response = upload(client, "empty.txt", b"")

    assert response.status_code == 400
    assert response.json()["error"] == "EMPTY_FILE"


def test_upload_malformed_docx(client, mock_embed):
    response = upload(client, "broken.docx", b"not a zip archive at all")

    assert response.status_code == 422
    assert response.json()["error"] == "EXTRACTION_FAILED"


def test_upload_pdf_bad_magic_bytes(client, mock_embed):
    response = upload(client, "fake.pdf", b"this is definitely not a pdf")

    assert response.status_code == 422
    assert response.json()["error"] == "PDF_CORRUPTED"


# ============================================================================
# Listing and deleting multi-format documents
# ============================================================================


def test_list_shows_all_formats(client, mock_embed):
    upload(client, "notes.txt", SAMPLES["notes.txt"]())
    upload(client, "data.csv", SAMPLES["data.csv"]())

    response = client.get("/documents/list")
    assert response.status_code == 200
    data = response.json()
    assert data["total_documents"] == 2
    types = {doc["file_type"] for doc in data["documents"]}
    assert types == {"txt", "csv"}


def test_legacy_chunks_without_file_type_still_list(client, mock_embed):
    """Documents stored before 6-C (no file_type metadata) must still list."""
    from backend.app.services.chroma_service import ChromaService

    service = ChromaService()
    service.add_chunks(
        [
            {
                "text": "legacy chunk",
                "page": 1,
                "file_name": "legacy.pdf",
                "chunk_index": 0,
                "doc_id": "legacy-doc-id",
                "upload_time": "2026-01-01T00:00:00Z",
                "size_bytes": 100,
            }
        ]
    )

    response = client.get("/documents/list")
    assert response.status_code == 200
    legacy = next(
        d for d in response.json()["documents"] if d["doc_id"] == "legacy-doc-id"
    )
    assert legacy["file_type"] is None  # absent metadata tolerated, not fabricated


def test_delete_one_of_multiple_formats(client, mock_embed):
    docx_resp = upload(client, "report.docx", SAMPLES["report.docx"]())
    md_resp = upload(client, "guide.md", SAMPLES["guide.md"]())
    docx_id = docx_resp.json()["doc_id"]

    response = client.delete(f"/documents/{docx_id}")
    assert response.status_code == 200
    assert response.json()["chunks_deleted"] >= 1

    remaining = client.get("/documents/list").json()
    assert remaining["total_documents"] == 1
    assert remaining["documents"][0]["file_name"] == "guide.md"


# ============================================================================
# Retrieval from chunks produced by each format (fully offline)
# ============================================================================


class StubEmbeddings:
    """Deterministic 768-dim embeddings without any network access.

    Vectors are derived from a hash of the input text so that different
    texts produce different embeddings (ChromaDB's HNSW index misbehaves
    when all stored vectors are identical).
    """

    def embed(self, text):
        import hashlib
        import numpy as np

        single = isinstance(text, str)
        texts = [text] if single else text
        vectors = []
        for t in texts:
            h = int(hashlib.md5(t.encode()).hexdigest(), 16) % (10**6)
            # Derive a pseudo-random but deterministic vector from the hash
            rng = np.random.RandomState(h % (2**31))
            vectors.append(rng.rand(768).astype(np.float32))
        matrix = np.array(vectors, dtype=np.float32)
        return matrix[0] if single else matrix


def test_retrieval_from_each_format(tmp_path):
    """Each format's chunks must be retrievable with format-distinct keywords."""
    from backend.app.services.chroma_service import ChromaService
    from backend.app.services.document_processor import process_document
    from backend.app.services.hybrid_retriever import HybridRetriever

    db_path = tempfile.mkdtemp(prefix="mf_retrieve_")
    try:
        service = ChromaService(db_path=db_path, embeddings_service=StubEmbeddings())

        queries = {
            "doc.pdf": ("doc.pdf", make_pdf_bytes("hybrid retrieval sample"), "hybrid"),

            "report.docx": (
                "report.docx",
                make_docx_bytes(
                    [("Installation", "Heading 1"), ("installation steps here.", "Normal")]
                ),
                "installation",
            ),
            "notes.txt": ("notes.txt", b"plainwords appear in plain documents.", "plainwords"),
            "guide.md": ("guide.md", b"# Markdown Guide\nmarkdown structure matters\n", "markdown"),
            "data.csv": ("data.csv", b"Region,Revenue\nEMEA,4200\n", "EMEA"),
            "book.xlsx": (
                "book.xlsx",
                make_xlsx_bytes({"Sales": (["Item"], [["revenuegadget"]])}),
                "revenuegadget",
            ),
        }

        for doc_id, (name, content, _) in queries.items():
            path = tmp_path / name
            path.write_bytes(content)
            chunks = process_document(str(path), name)
            for chunk in chunks:
                chunk["doc_id"] = doc_id
                chunk["upload_time"] = "2026-01-01T00:00:00Z"
                chunk["size_bytes"] = len(content)
            service.add_chunks(chunks)

        retriever = HybridRetriever(chroma_service=service)
        for file_name, (_, _, keyword) in queries.items():
            results = retriever.retrieve(keyword, n_results=10)
            files_in_results = {m["file_name"] for m in results["metadatas"]}
            assert file_name in files_in_results, (
                f"{file_name} not retrieved for keyword '{keyword}'"
            )

        # Source attribution: PDF chunks cite pages; tabular chunks cite rows
        csv_results = retriever.retrieve("EMEA", n_results=10)
        csv_meta = next(
            m for m in csv_results["metadatas"] if m["file_name"] == "data.csv"
        )
        assert csv_meta.get("rows")
        assert csv_meta.get("page") is None
    finally:
        shutil.rmtree(db_path, ignore_errors=True)